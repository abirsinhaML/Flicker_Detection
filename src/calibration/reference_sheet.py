"""Extraction of the reference label set from the delivery spreadsheet.

The reference labels exist only as cell fill colours, which no CSV or dataframe
export preserves, so the workbook has to be read as a workbook.  Colours are
matched exactly rather than by proximity: the sheet uses three specific fills,
and silently mapping an unrecognised colour onto the nearest label would invent
ground truth.
"""

from __future__ import annotations

import csv
import logging
from collections.abc import Iterator
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

# Red = extreme, orange = mild, green = clean.  ARGB as written by Sheets.
FILL_LABELS: dict[str, str] = {
    "FFF4CCCC": "extreme",
    "FFFCE5CD": "mild",
    "FFD9EAD3": "none",
}
LABEL_FIELDS: tuple[str, ...] = ("video_key", "label")
_KEY_COLUMNS = ("key", "video_key")


def extract_labels(workbook_path: str | Path) -> list[tuple[str, str]]:
    """Return ``(video_key, label)`` for every colour-coded row, in sheet order.

    A row carrying more than one distinct fill colour is ambiguous and is
    skipped with a warning rather than resolved by guesswork.
    """
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Reference workbook not found: {workbook_path}")

    workbook = openpyxl.load_workbook(workbook_path)
    labels: list[tuple[str, str]] = []
    seen: set[str] = set()
    for sheet in workbook.worksheets:
        key_column = _find_key_column(sheet)
        for row in sheet.iter_rows(min_row=2):
            key = _cell_text(row[key_column]) if key_column < len(row) else None
            if not key:
                continue
            found = {label for label in map(_fill_label, row) if label}
            if not found:
                continue
            if len(found) > 1:
                logger.warning("Skipping %s: conflicting fill colours %s", key, sorted(found))
                continue
            if key in seen:
                logger.warning("Skipping duplicate key %s", key)
                continue
            seen.add(key)
            labels.append((key, found.pop()))
    return labels


def write_labels(labels: list[tuple[str, str]], output_path: str | Path) -> int:
    """Write labels as the ``video_key,label`` CSV the calibrator consumes."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=LABEL_FIELDS)
        writer.writeheader()
        for key, label in labels:
            writer.writerow({"video_key": key, "label": label})
    return len(labels)


def _fill_label(cell: object) -> str | None:
    fill = getattr(cell, "fill", None)
    if fill is None or fill.patternType != "solid":
        return None
    rgb = getattr(fill.fgColor, "rgb", None)
    return FILL_LABELS.get(str(rgb).upper()) if isinstance(rgb, str) else None


def _find_key_column(sheet: object) -> int:
    """Locate the column holding video keys, defaulting to the first."""
    for row in _header_rows(sheet):
        for index, cell in enumerate(row):
            text = _cell_text(cell)
            if text and text.strip().lower() in _KEY_COLUMNS:
                return index
        break
    return 0


def _header_rows(sheet: object) -> Iterator[tuple[object, ...]]:
    yield from sheet.iter_rows(min_row=1, max_row=1)  # type: ignore[attr-defined]


def _cell_text(cell: object) -> str | None:
    value = getattr(cell, "value", None)
    if value is None:
        return None
    text = str(value).strip()
    return text or None
