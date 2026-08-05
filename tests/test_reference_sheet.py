"""Tests for reading reference labels out of spreadsheet cell colours."""

from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from src.calibration.reference_sheet import LABEL_FIELDS, extract_labels, write_labels

EXTREME = "FFF4CCCC"
MILD = "FFFCE5CD"
CLEAN = "FFD9EAD3"
UNKNOWN = "FF0000FF"


def _workbook(rows: list[tuple[str, str | None]], path: Path, header: str = "key") -> None:
    """Write a sheet whose rows carry a fill colour in their second column."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([header, "size_bytes"])
    for key, colour in rows:
        sheet.append([key, 1])
        if colour:
            sheet.cell(row=sheet.max_row, column=2).fill = PatternFill(
                patternType="solid", fgColor=colour
            )
    workbook.save(path)


class ExtractLabelsTests(unittest.TestCase):
    def test_maps_the_three_reference_colours(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sheet.xlsx"
            _workbook(
                [("a.mp4", EXTREME), ("b.mp4", MILD), ("c.mp4", CLEAN)],
                path,
            )
            labels = extract_labels(path)

        self.assertEqual(
            labels,
            [("a.mp4", "extreme"), ("b.mp4", "mild"), ("c.mp4", "none")],
        )

    def test_uncoloured_rows_are_not_labelled(self) -> None:
        """Absence of a colour is absence of a label, not a clean verdict."""
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sheet.xlsx"
            _workbook([("a.mp4", None), ("b.mp4", MILD)], path)
            labels = extract_labels(path)

        self.assertEqual(labels, [("b.mp4", "mild")])

    def test_unrecognised_colours_are_ignored_rather_than_guessed(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sheet.xlsx"
            _workbook([("a.mp4", UNKNOWN)], path)

            self.assertEqual(extract_labels(path), [])

    def test_conflicting_colours_in_one_row_are_skipped_with_a_warning(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sheet.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["key", "a", "b"])
            sheet.append(["a.mp4", 1, 2])
            sheet.cell(row=2, column=2).fill = PatternFill(patternType="solid", fgColor=MILD)
            sheet.cell(row=2, column=3).fill = PatternFill(patternType="solid", fgColor=EXTREME)
            workbook.save(path)

            with self.assertLogs("src.calibration.reference_sheet", level="WARNING"):
                self.assertEqual(extract_labels(path), [])

    def test_finds_the_key_column_by_header_name(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sheet.xlsx"
            _workbook([("a.mp4", MILD)], path, header="video_key")

            self.assertEqual(extract_labels(path), [("a.mp4", "mild")])

    def test_missing_workbook_is_an_error(self) -> None:
        with self.assertRaises(FileNotFoundError):
            extract_labels("/nonexistent/sheet.xlsx")


class WriteLabelsTests(unittest.TestCase):
    def test_writes_the_schema_the_calibrator_reads(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "nested" / "labels.csv"
            written = write_labels([("a.mp4", "mild"), ("b.mp4", "none")], path)
            with path.open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))

        self.assertEqual(written, 2)
        self.assertEqual(tuple(rows[0]), LABEL_FIELDS)
        self.assertEqual(rows[0], {"video_key": "a.mp4", "label": "mild"})


class CommittedLabelsTests(unittest.TestCase):
    """Guard the checked-in label set, which calibration depends on."""

    def test_reference_labels_are_present_and_well_formed(self) -> None:
        path = Path("data/reference_labels.csv")
        if not path.exists():
            self.skipTest("reference labels have not been extracted")
        with path.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))

        self.assertEqual(tuple(rows[0]), LABEL_FIELDS)
        self.assertTrue(all(row["label"] in {"none", "mild", "extreme"} for row in rows))
        self.assertEqual(len({row["video_key"] for row in rows}), len(rows))
        # Every band must be represented or thresholds cannot be fitted.
        self.assertEqual({row["label"] for row in rows}, {"none", "mild", "extreme"})


if __name__ == "__main__":
    unittest.main()
